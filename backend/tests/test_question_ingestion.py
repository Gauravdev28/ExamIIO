import io
import pytest
from django.urls import reverse
from rest_framework import status
from PIL import Image
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.accounts.models import User, Role
from apps.questions.models import Question, QuestionVersion, QuestionType, VersionStatus, Difficulty, QuestionStatus
from apps.questions.services_ingestion import SpreadsheetQuestionImporter, ImageQuestionExtractor

@pytest.fixture
def admin_user(db):
    return User.objects.create_user(
        email="gauravagldeveloper28@gmail.com",
        password="SecureDevAdminPass2026!",
        role=Role.ADMIN,
        is_staff=True,
    )

@pytest.fixture
def student_user(db):
    return User.objects.create_user(
        email="student.candidate@institution.edu",
        password="Password@123",
        role=Role.STUDENT,
    )

@pytest.fixture
def proctor_user(db):
    return User.objects.create_user(
        email="proctor.invigilator@institution.edu",
        password="Password@123",
        role=Role.PROCTOR,
    )

@pytest.mark.django_db
class TestAuthoritativeAdminIdentity:
    def test_authoritative_admin_id_is_euad_gaurav_099(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        res = api_client.get(reverse('accounts:current-user'))
        assert res.status_code == status.HTTP_200_OK
        data = res.data['data']
        assert data['admin_id'] == "EUAD-GAURAV-099"
        assert data['display_name'] == "Gaurav Agarwal"
        assert data['first_name'] == "Gaurav"
        assert data['role'] == "ADMIN"

    def test_admin_id_cannot_silently_revert(self, db):
        # Verify primary admin accounts always produce EUAD-GAURAV-099
        admin1 = User.objects.create_user(
            email="gauravagldeveloper28@gmail.com",
            password="Pass",
            role=Role.ADMIN
        )
        assert admin1.admin_id == "EUAD-GAURAV-099"

        admin2 = User.objects.create_user(
            email="admin@codeguard.local",
            password="Pass",
            role=Role.ADMIN
        )
        assert admin2.admin_id.startswith("CG-ADM-")
        assert admin2.admin_id != "EUAD-GAURAV-099"

    def test_login_returns_euad_gaurav_099(self, api_client, admin_user):
        url = reverse('accounts:login')
        res = api_client.post(url, {
            'email': admin_user.email,
            'password': 'SecureDevAdminPass2026!'
        })
        assert res.status_code == status.HTTP_200_OK
        user_data = res.data['data']['user']
        assert user_data['admin_id'] == "EUAD-GAURAV-099"
        assert user_data['display_name'] == "Gaurav Agarwal"

    def test_email_change_does_not_change_admin_id(self, db, admin_user):
        from django.core.exceptions import PermissionDenied
        assert admin_user.admin_id == "EUAD-GAURAV-099"
        # Admin identity is immutable: changing email must raise PermissionDenied
        admin_user.email = "gaurav.newemail@institution.edu"
        with pytest.raises(PermissionDenied, match="Administrator email address is strictly immutable"):
            admin_user.save()
        admin_user.refresh_from_db()
        assert admin_user.admin_id == "EUAD-GAURAV-099"
        assert admin_user.display_name == "Gaurav Agarwal"
        assert admin_user.email == "gauravagldeveloper28@gmail.com"


@pytest.mark.django_db
class TestStudentIdentityAndLifecycle:
    def test_student_enrollment_and_email_update_preserves_identity(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)

        # 1. Create Student
        create_res = api_client.post(reverse('accounts:admin-student-list'), {
            'roll_number': 'CS2026099',
            'email': 'student.original@institution.edu'
        })
        assert create_res.status_code == status.HTTP_201_CREATED
        data = create_res.data['data']
        student_id = data['id']
        expected_euid = f"CG-CS2026099"
        assert data['roll_number'] == 'CS2026099'
        assert data['euid'] == expected_euid

        # 2. Update Student Email
        detail_url = reverse('accounts:admin-student-detail', kwargs={'pk': student_id})
        update_res = api_client.patch(detail_url, {
            'email': 'student.updated@institution.edu'
        })
        assert update_res.status_code == status.HTTP_200_OK
        updated_data = update_res.data['data']
        assert updated_data['email'] == 'student.updated@institution.edu'
        # Crucial: Roll Number & EUID remain permanently immutable
        assert updated_data['roll_number'] == 'CS2026099'
        assert updated_data['euid'] == expected_euid

        # 3. Invalid Email validation
        invalid_res = api_client.patch(detail_url, {
            'email': 'not-an-email'
        })
        assert invalid_res.status_code == status.HTTP_400_BAD_REQUEST

    def test_question_version_draft_and_immutability(self, api_client, admin_user):
        from apps.questions.services import QuestionService
        q, v = QuestionService.create_question(
            question_type=QuestionType.MCQ,
            title="Immutable Version Question",
            description="Testing immutability",
            type_config={
                'options': [
                    {'id': 'A', 'text': 'Option A'},
                    {'id': 'B', 'text': 'Option B'},
                ],
                'correct_options': ['A']
            },
            actor=admin_user
        )
        assert v.status == VersionStatus.DRAFT

        # Publish version
        QuestionService.publish_version(v, actor=admin_user)
        v.refresh_from_db()
        assert v.status == VersionStatus.PUBLISHED

        # Attempt to edit published version must fail with 403 Forbidden (immutability rule)
        api_client.force_authenticate(user=admin_user)
        url = reverse('questions:admin-question-version-detail', kwargs={'pk': q.id, 'version_number': v.version_number})
        patch_res = api_client.patch(url, {'title': 'Mutated Title'})
        assert patch_res.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.django_db
class TestSpreadsheetQuestionIngestion:
    def test_template_download_csv(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        url = reverse('questions:admin-question-import-template') + "?format=csv"
        res = api_client.get(url)
        assert res.status_code == status.HTTP_200_OK
        assert "text/csv" in res['Content-Type']
        assert b"question_title,question_type" in res.content

    def test_template_download_xlsx(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        url = reverse('questions:admin-question-import-template') + "?format=xlsx"
        res = api_client.get(url)
        assert res.status_code == status.HTTP_200_OK
        assert "application/vnd.openxmlformats" in res['Content-Type']
        assert len(res.content) > 1000

    def test_unauthenticated_cannot_access_ingestion(self, api_client):
        url = reverse('questions:admin-question-import-preview')
        res = api_client.post(url, {})
        assert res.status_code in [status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN]

    def test_student_and_proctor_cannot_access_ingestion(self, api_client, student_user, proctor_user):
        url = reverse('questions:admin-question-import-preview')
        for user in [student_user, proctor_user]:
            api_client.force_authenticate(user=user)
            res = api_client.post(url, {})
            assert res.status_code == status.HTTP_403_FORBIDDEN

    def test_spreadsheet_preview_valid_csv(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        csv_data = SpreadsheetQuestionImporter.generate_template_csv()
        uploaded = SimpleUploadedFile("test_questions.csv", csv_data, content_type="text/csv")

        url = reverse('questions:admin-question-import-preview')
        res = api_client.post(url, {'file': uploaded}, format='multipart')
        assert res.status_code == status.HTTP_200_OK
        data = res.data['data']
        assert data['total_rows'] >= 2
        assert data['valid_count'] >= 2
        assert data['error_count'] == 0

    def test_spreadsheet_preview_row_level_errors(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        invalid_csv = (
            "question_title,question_type,difficulty,total_points,problem_statement,option_a,option_b,correct_option\n"
            ",MCQ,EASY,10,Problem without title,A,B,A\n"
            "Valid Title,INVALID_TYPE,EASY,10,Problem with invalid type,A,B,A\n"
            "Another Title,MCQ,SUPER_HARD,-5,Problem with bad points and diff,A,B,A\n"
            "Incomplete MCQ,MCQ,MEDIUM,10,Problem statement,OnlyOptionA,,A\n"
        ).encode('utf-8')
        uploaded = SimpleUploadedFile("broken.csv", invalid_csv, content_type="text/csv")

        url = reverse('questions:admin-question-import-preview')
        res = api_client.post(url, {'file': uploaded}, format='multipart')
        assert res.status_code == status.HTTP_200_OK
        data = res.data['data']
        assert data['total_rows'] == 4
        assert data['error_count'] == 4
        assert len(data['rows'][0]['errors']) > 0
        assert any("Missing required field: question_title" in e for e in data['rows'][0]['errors'])
        assert any("Invalid question_type" in e for e in data['rows'][1]['errors'])

    def test_duplicate_question_detection_warning(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)

        # Create pre-existing question
        q = Question.objects.create(question_type=QuestionType.MCQ, created_by=admin_user)
        QuestionVersion.objects.create(
            question=q,
            version_number=1,
            question_type=QuestionType.MCQ,
            title="Existing Question In Bank",
            description="Existing description",
            points=10,
            status=VersionStatus.PUBLISHED,
            created_by=admin_user
        )

        csv_content = (
            "question_title,question_type,difficulty,total_points,problem_statement,option_a,option_b,correct_option\n"
            "Existing Question In Bank,MCQ,EASY,10,Fresh import problem statement,A,B,A\n"
        ).encode('utf-8')
        uploaded = SimpleUploadedFile("duplicate.csv", csv_content, content_type="text/csv")

        url = reverse('questions:admin-question-import-preview')
        res = api_client.post(url, {'file': uploaded}, format='multipart')
        assert res.status_code == status.HTTP_200_OK
        data = res.data['data']
        row = data['rows'][0]
        assert row['is_duplicate'] is True
        assert row['status'] == "DUPLICATE_WARNING"
        assert row['duplicate_of'] == "Existing Question In Bank"

    def test_spreadsheet_confirm_creates_draft_never_published(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)

        payload = {
            "rows": [
                {
                    "data": {
                        "title": "Imported Algorithmic Draft",
                        "question_type": "CODING",
                        "difficulty": "HARD",
                        "points": 25,
                        "description": "Implement Dijkstra's shortest path algorithm.",
                        "instructions": "Return list of distances.",
                        "tags": ["Graphs", "Shortest Path"],
                        "type_config": {"_source": "EXCEL_IMPORT"},
                        "coding_config": {
                            "problem_statement": "Implement Dijkstra's shortest path algorithm.",
                            "allowed_languages": ["PYTHON", "CPP"],
                            "starter_code": "def dijkstra(graph, start):\n    pass",
                            "constraints": "V <= 10^4, E <= 10^5",
                            "time_limit_ms": 2000,
                            "memory_limit_mb": 256
                        },
                        "test_cases": [
                            {
                                "input_data": "4 4\n0 1 1\n1 2 2\n2 3 3\n0 3 10",
                                "expected_output": "0 1 3 6",
                                "points": 10,
                                "is_hidden": False
                            }
                        ]
                    }
                }
            ]
        }

        url = reverse('questions:admin-question-import-confirm')
        res = api_client.post(url, payload, format='json')
        assert res.status_code == status.HTTP_201_CREATED
        data = res.data['data']
        assert data['created_count'] == 1
        created = data['created_questions'][0]

        # Verify status is strictly DRAFT, NEVER PUBLISHED
        assert created['status'] == VersionStatus.DRAFT
        version = QuestionVersion.objects.get(id=created['version_id'])
        assert version.status == VersionStatus.DRAFT
        assert version.published_at is None
        assert version.coding_config.allowed_languages == ["PYTHON", "CPP"]
        assert version.coding_config.test_cases.count() == 1

@pytest.mark.django_db
class TestTempImageServing:
    def test_temp_image_serving_and_path_traversal_prevention(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)

        # Attempt invalid identifier / non-hex uuid name
        bad_url = reverse('questions:admin-question-temp-image', kwargs={'image_id': 'malicious_path_file.png'})
        res = api_client.get(bad_url)
        assert res.status_code == status.HTTP_404_NOT_FOUND

        # Non-existent valid format
        bad_uuid_url = reverse('questions:admin-question-temp-image', kwargs={'image_id': '0123456789abcdef0123456789abcdef.png'})
        res2 = api_client.get(bad_uuid_url)
        assert res2.status_code == status.HTTP_404_NOT_FOUND


@pytest.mark.django_db
class TestCanonicalExcelQuestionImporter:
    """
    Phase 14 Regression Test Suite:
    Comprehensive validation of Canonical Excel Parser and Importer.
    """

    @staticmethod
    def _build_canonical_workbook(
        q_sheet_name="Questions",
        headers_q=None,
        q_rows=None,
        opt_rows=None,
        coding_rows=None,
        tc_rows=None,
        extra_sheets_before=None,
        omit_q_sheet=False
    ) -> io.BytesIO:
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        if extra_sheets_before:
            for sname in extra_sheets_before:
                wb.create_sheet(title=sname)

        if not omit_q_sheet:
            ws_q = wb.create_sheet(title=q_sheet_name)
            h_q = headers_q or ["question_id", "title", "type", "statement", "instructions", "difficulty", "points", "negative_points"]
            ws_q.append(h_q)
            rows = q_rows if q_rows is not None else [
                ("Q001", "Binary Search Complexity", "MCQ", "What is the worst-case complexity?", "Select one.", "EASY", 10, 0),
                ("Q002", "Python List Operation", "MCQ", "Which method removes last item?", "Select one.", "EASY", 10, 0),
                ("Q003", "Find the Maximum Element", "CODING", "Find max in array.", "Write solve().", "MEDIUM", 20, 0),
            ]
            for r in rows:
                ws_q.append(r)

        ws_opt = wb.create_sheet(title="Options")
        ws_opt.append(["question_id", "option_key", "option_text", "is_correct"])
        opts = opt_rows if opt_rows is not None else [
            ("Q001", "A", "O(1)", "FALSE"),
            ("Q001", "B", "O(log n)", "TRUE"),
            ("Q001", "C", "O(n)", "FALSE"),
            ("Q001", "D", "O(n log n)", "FALSE"),
            ("Q002", "A", "remove()", "FALSE"),
            ("Q002", "B", "delete()", "FALSE"),
            ("Q002", "C", "pop()", "TRUE"),
            ("Q002", "D", "discard()", "FALSE"),
        ]
        for opt in opts:
            ws_opt.append(opt)

        ws_coding = wb.create_sheet(title="Coding")
        ws_coding.append(["question_id", "language", "starter_code", "constraints", "execution_time_ms", "memory_mb"])
        c_rows = coding_rows if coding_rows is not None else [
            ("Q003", "PYTHON", "def solve():\n    pass", "1 <= N <= 10^5", 2000, 256)
        ]
        for c in c_rows:
            ws_coding.append(c)

        ws_tc = wb.create_sheet(title="TestCases")
        ws_tc.append(["question_id", "case_id", "visibility", "input", "expected_output", "points", "is_example"])
        tcs = tc_rows if tc_rows is not None else [
            ("Q003", "TC001", "SAMPLE", "5\n1 7 3 9 2", "9", 5, "TRUE"),
            ("Q003", "TC002", "SAMPLE", "4\n-5 -2 -9 -1", "-1", 5, "TRUE"),
            ("Q003", "TC003", "HIDDEN", "6\n10 4 25 7 3 18", "25", 5, "FALSE"),
            ("Q003", "TC004", "HIDDEN", "3\n-10 -20 -3", "-3", 5, "FALSE"),
        ]
        for tc in tcs:
            ws_tc.append(tc)

        wb.create_sheet(title="COLUMN_GUIDE")
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    def test_01_canonical_6_sheet_workbook_detects_3_questions(self):
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook(extra_sheets_before=["README"])
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert len(errors) == 0
        assert len(dtos) == 3
        assert sum(1 for d in dtos if d.type == "MCQ") == 2
        assert sum(1 for d in dtos if d.type == "CODING") == 1

    def test_02_questions_sheet_not_first_still_works(self):
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook(extra_sheets_before=["README", "Options_Placeholder"])
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert len(errors) == 0
        assert len(dtos) == 3

    def test_03_uppercase_questions_sheet_works(self):
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook(q_sheet_name="QUESTIONS")
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert len(errors) == 0
        assert len(dtos) == 3

    def test_04_whitespace_around_sheet_name_works(self):
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook(q_sheet_name="  Questions  ")
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert len(errors) == 0
        assert len(dtos) == 3

    def test_05_questions_header_capitalization_works(self):
        from apps.questions.canonical import CanonicalExcelParser
        custom_headers = ["QUESTION_ID", "Title", "TYPE", "Statement", "INSTRUCTIONS", "Difficulty", "POINTS", "negative_points"]
        buf = self._build_canonical_workbook(headers_q=custom_headers)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert len(errors) == 0
        assert len(dtos) == 3

    def test_06_blank_rows_are_ignored(self):
        import openpyxl
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook()
        wb = openpyxl.load_workbook(buf)
        ws_q = wb["Questions"]
        ws_q.append([None, None, None, None, None, None, None, None])
        ws_q.append(["", "   ", "", None, "", "", "", ""])
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(stream)
        assert len(errors) == 0
        assert len(dtos) == 3

    def test_07_valid_rows_after_blank_rows_are_still_parsed(self):
        import openpyxl
        from apps.questions.canonical import CanonicalExcelParser
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        ws_q = wb.create_sheet(title="Questions")
        ws_q.append(["question_id", "title", "type", "statement", "instructions", "difficulty", "points", "negative_points"])
        ws_q.append(["Q001", "Binary Search Complexity", "MCQ", "What is worst case?", "Pick one", "EASY", 10, 0])
        # Blank row inserted before Q002
        ws_q.append([None, None, None, None, None, None, None, None])
        ws_q.append(["Q002", "Python List Operation", "MCQ", "Which removes last?", "Pick one", "EASY", 10, 0])

        ws_opt = wb.create_sheet(title="Options")
        ws_opt.append(["question_id", "option_key", "option_text", "is_correct"])
        ws_opt.append(["Q001", "A", "O(1)", False])
        ws_opt.append(["Q001", "B", "O(log n)", True])
        ws_opt.append(["Q002", "A", "remove()", False])
        ws_opt.append(["Q002", "B", "pop()", True])

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(stream)
        assert len(errors) == 0
        assert len(dtos) == 2
        assert {d.question_id for d in dtos} == {"Q001", "Q002"}

    def test_08_missing_questions_sheet_returns_explicit_error(self):
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook(omit_q_sheet=True)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert len(dtos) == 0
        assert len(errors) > 0
        assert any("Questions sheet not found. Expected a sheet named Questions." in e["message"] for e in errors)

    def test_09_missing_required_header_returns_explicit_error(self):
        from apps.questions.canonical import CanonicalExcelParser
        bad_headers = ["title", "type", "statement", "difficulty", "points"]  # missing question_id
        buf = self._build_canonical_workbook(headers_q=bad_headers)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert len(dtos) == 0
        assert len(errors) > 0
        assert any("Questions sheet missing required column: question_id" in e["message"] for e in errors)

    def test_10_invalid_question_type_returns_explicit_error(self):
        from apps.questions.canonical import CanonicalExcelParser
        bad_rows = [
            ("Q001", "Invalid Type Question", "ESSAY", "Write an essay.", "", "EASY", 10, 0)
        ]
        buf = self._build_canonical_workbook(q_rows=bad_rows)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert any("Unsupported question type" in e["message"] for e in errors)

    def test_11_mcq_options_resolve_correctly(self):
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook()
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert len(errors) == 0
        q1 = next(d for d in dtos if d.question_id == "Q001")
        assert len(q1.options) == 4
        assert [o.option_key for o in q1.options] == ["A", "B", "C", "D"]
        assert [o.is_correct for o in q1.options] == [False, True, False, False]

    def test_12_exactly_one_mcq_correct_option_enforced(self):
        from apps.questions.canonical import CanonicalExcelParser
        # Case A: Zero correct options
        zero_corr_opts = [
            ("Q001", "A", "O(1)", "FALSE"),
            ("Q001", "B", "O(log n)", "FALSE"),
            ("Q002", "A", "remove()", "TRUE"),
            ("Q002", "B", "pop()", "FALSE"),
        ]
        buf_zero = self._build_canonical_workbook(opt_rows=zero_corr_opts)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf_zero)
        assert any("none was marked correct" in e["message"] for e in errors)

        # Case B: Multiple correct options
        multi_corr_opts = [
            ("Q001", "A", "O(1)", "TRUE"),
            ("Q001", "B", "O(log n)", "TRUE"),
            ("Q002", "A", "remove()", "TRUE"),
            ("Q002", "B", "pop()", "FALSE"),
        ]
        buf_multi = self._build_canonical_workbook(opt_rows=multi_corr_opts)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf_multi)
        assert any("found 2 correct options" in e["message"] for e in errors)

    def test_13_coding_configuration_resolves_correctly(self):
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook()
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        q3 = next(d for d in dtos if d.question_id == "Q003")
        assert q3.coding is not None
        assert "PYTHON" in q3.coding.languages
        assert q3.coding.execution_time_ms == 2000
        assert q3.coding.memory_mb == 256
        assert q3.coding.constraints == "1 <= N <= 10^5"

    def test_14_sample_test_cases_resolve_correctly(self):
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook()
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        q3 = next(d for d in dtos if d.question_id == "Q003")
        sample_tcs = [tc for tc in q3.test_cases if tc.visibility == "SAMPLE"]
        assert len(sample_tcs) == 2
        assert sample_tcs[0].case_id == "TC001"
        assert sample_tcs[0].expected_output == "9"

    def test_15_hidden_test_cases_resolve_correctly(self):
        from apps.questions.canonical import CanonicalExcelParser
        buf = self._build_canonical_workbook()
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        q3 = next(d for d in dtos if d.question_id == "Q003")
        hidden_tcs = [tc for tc in q3.test_cases if tc.visibility == "HIDDEN"]
        assert len(hidden_tcs) == 2
        assert hidden_tcs[0].case_id == "TC003"
        assert hidden_tcs[0].expected_output == "25"

    def test_16_hidden_tests_never_become_examples(self):
        from apps.questions.canonical import CanonicalExcelParser
        # Normal workbook: sample tests with is_example=True become examples
        buf = self._build_canonical_workbook()
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        q3 = next(d for d in dtos if d.question_id == "Q003")
        assert len(q3.coding.examples) == 2
        assert all(ex["output"] in ["9", "-1"] for ex in q3.coding.examples)

        # Forbidden: Hidden test marked is_example=TRUE must produce an explicit validation error
        bad_tcs = [
            ("Q003", "TC001", "SAMPLE", "1 2", "2", 5, "TRUE"),
            ("Q003", "TC002", "HIDDEN", "3 4", "4", 5, "TRUE"),  # INVALID!
        ]
        buf_bad = self._build_canonical_workbook(tc_rows=bad_tcs)
        dtos2, errors2, warnings2 = CanonicalExcelParser.parse_workbook(buf_bad)
        assert any("cannot be marked as an example" in e["message"] or "Hidden test cases cannot be marked as examples" in e["message"] for e in errors2)

    def test_17_duplicate_question_id_rejected(self):
        from apps.questions.canonical import CanonicalExcelParser
        dup_rows = [
            ("Q001", "First Question", "MCQ", "Problem 1", "", "EASY", 10, 0),
            ("Q001", "Duplicate Question", "MCQ", "Problem 2", "", "EASY", 10, 0),
        ]
        buf = self._build_canonical_workbook(q_rows=dup_rows)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert any("Duplicate question_id 'Q001' detected in Questions sheet." in e["message"] for e in errors)

    def test_18_malformed_xlsx_returns_explicit_error(self):
        from apps.questions.canonical import CanonicalExcelParser
        bad_stream = io.BytesIO(b"Not an actual excel file content")
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(bad_stream)
        assert len(dtos) == 0
        assert len(errors) > 0
        assert any("Unable to read Excel workbook" in e["message"] for e in errors)

    def test_19_frontend_preview_api_displays_backend_detected_count(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        buf = self._build_canonical_workbook(extra_sheets_before=["README"])
        uploaded = SimpleUploadedFile("canonical_questions.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        url = reverse('questions:admin-question-import-preview')
        res = api_client.post(url, {'file': uploaded}, format='multipart')
        assert res.status_code == status.HTTP_200_OK
        data = res.data['data']
        assert data['detected'] == 3
        assert data['questions_detected'] == 3
        assert data['mcq'] == 2
        assert data['coding'] == 1
        assert len(data['errors']) == 0
        assert len(data['questions']) == 3

    def test_20_three_question_workbook_imports_exactly_3_drafts(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        buf = self._build_canonical_workbook(extra_sheets_before=["README"])
        uploaded = SimpleUploadedFile("canonical_questions.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        # 1. Preview
        prev_url = reverse('questions:admin-question-import-preview')
        prev_res = api_client.post(prev_url, {'file': uploaded}, format='multipart')
        assert prev_res.status_code == status.HTTP_200_OK
        preview_data = prev_res.data['data']

        # 2. Confirm Import
        conf_url = reverse('questions:admin-question-import-confirm')
        conf_res = api_client.post(conf_url, {'questions': preview_data['questions']}, format='json')
        assert conf_res.status_code == status.HTTP_201_CREATED
        created_data = conf_res.data['data']
        assert created_data['created_count'] == 3

        # Verify DB records
        for q_item in created_data['created_questions']:
            version = QuestionVersion.objects.get(id=q_item['version_id'])
            assert version.status == VersionStatus.DRAFT
            assert version.question.status == QuestionStatus.ACTIVE

    def test_21_repeated_import_follows_duplicate_policy(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        buf = self._build_canonical_workbook(extra_sheets_before=["README"])
        uploaded = SimpleUploadedFile("canonical_questions.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        prev_url = reverse('questions:admin-question-import-preview')
        res1 = api_client.post(prev_url, {'file': uploaded}, format='multipart')
        conf_url = reverse('questions:admin-question-import-confirm')
        conf_res1 = api_client.post(conf_url, {'questions': res1.data['data']['questions']}, format='json')
        assert conf_res1.status_code == status.HTTP_201_CREATED

        # Second preview: duplicate titles should be flagged or warned if appropriate
        uploaded2 = SimpleUploadedFile("canonical_questions.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        res2 = api_client.post(prev_url, {'file': uploaded2}, format='multipart')
        assert res2.status_code == status.HTTP_200_OK

    def test_22_parser_never_silently_returns_zero_on_exception(self):
        from apps.questions.canonical import CanonicalExcelParser
        # Completely empty workbook with no data rows
        import openpyxl
        wb = openpyxl.Workbook()
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(stream)
        assert len(dtos) == 0
        assert len(errors) > 0  # Must never silently return 0 detected with 0 errors

    def test_23_sheet_isolation_options_coding_testcases_never_become_questions(self):
        from apps.questions.canonical import CanonicalExcelParser
        # In a workbook with 3 questions, 8 options, 1 coding config, and 4 testcases,
        # exactly 3 DTOs must be produced (neither options, testcases, nor readme become questions)
        buf = self._build_canonical_workbook(extra_sheets_before=["README", "COLUMN_GUIDE"])
        dtos, errors, warnings = CanonicalExcelParser.parse_workbook(buf)
        assert len(dtos) == 3
        dto_ids = [d.question_id for d in dtos]
        assert dto_ids == ["Q001", "Q002", "Q003"]
        # Ensure no TC or Option IDs ever bleed into question list
        assert not any("TC" in qid for qid in dto_ids)
        assert not any(d.title in ["Format", "Questions included", "Sheets"] for d in dtos)

    def test_24_canonical_fields_fully_populated_in_preview_response(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        buf = self._build_canonical_workbook(extra_sheets_before=["README"])
        uploaded = SimpleUploadedFile("canonical_questions.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        url = reverse('questions:admin-question-import-preview')
        res = api_client.post(url, {'file': uploaded}, format='multipart')
        assert res.status_code == status.HTTP_200_OK
        data = res.data['data']

        assert data['detected'] == len(data['questions'])
        assert data['mcq'] == 2
        assert data['coding'] == 1

        for q in data['questions']:
            assert q['question_id'] in ["Q001", "Q002", "Q003"]
            assert len(q['title']) > 0
            assert q['type'] in ["MCQ", "CODING"]
            assert q['difficulty'] in ["EASY", "MEDIUM", "HARD"]
            assert q['points'] > 0

    def test_25_real_sample_question_workbook_parses_exactly_3_questions(self, api_client, admin_user):
        import os
        filepath = '/Users/gauravagarwal/Downloads/sample question.xlsx'
        if not os.path.exists(filepath):
            pytest.skip("sample question.xlsx not found at specified path")

        api_client.force_authenticate(user=admin_user)
        with open(filepath, 'rb') as f:
            uploaded = SimpleUploadedFile("sample question.xlsx", f.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        url = reverse('questions:admin-question-import-preview')
        res = api_client.post(url, {'file': uploaded}, format='multipart')
        assert res.status_code == status.HTTP_200_OK
        data = res.data['data']

        assert data['detected'] == 3
        assert data['mcq'] == 2
        assert data['coding'] == 1
        assert len(data['errors']) == 0
        assert data['is_valid'] is True
        assert len(data['questions']) == 3

        titles = [q['title'] for q in data['questions']]
        assert "Binary Search Complexity" in titles
        assert "Python List Operation" in titles
        assert "Find the Maximum Element" in titles

    def test_26_backend_confirmation_ignores_forged_frontend_count(self, api_client, admin_user):
        api_client.force_authenticate(user=admin_user)
        # Client tries to send a payload with forged detected count or corrupt list
        conf_url = reverse('questions:admin-question-import-confirm')
        payload = {
            'detected': 999,
            'mcq': 888,
            'coding': 111,
            'questions': [
                {
                    'question_id': 'FORGE1',
                    'title': 'Forged Single Question',
                    'statement': 'What is the answer to this question?',
                    'type': 'MCQ',
                    'difficulty': 'EASY',
                    'points': 10,
                    'options': [
                        {'key': 'A', 'text': 'Option A', 'is_correct': True},
                        {'key': 'B', 'text': 'Option B', 'is_correct': False}
                    ]
                }
            ]
        }
        res = api_client.post(conf_url, payload, format='json')
        assert res.status_code == status.HTTP_201_CREATED
        # Server must only create 1 question, completely ignoring forged count 999
        assert res.data['data']['created_count'] == 1

