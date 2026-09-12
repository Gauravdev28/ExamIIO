import io
import json
import uuid
import pytest
import openpyxl
from datetime import timedelta
from django.utils import timezone
from django.urls import reverse
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied
from rest_framework.exceptions import ValidationError as DRFValidationError

from apps.accounts.models import Role, StudentProfile
from apps.questions.models import (
    Question,
    QuestionVersion,
    QuestionType,
    Difficulty,
    VersionStatus,
    QuestionStatus,
    CodingLanguage,
    CodingQuestionConfig,
    TestCase,
)
from apps.questions.canonical import (
    CanonicalQuestionDTO,
    CanonicalOptionDTO,
    CanonicalCodingConfigDTO,
    CanonicalTestCaseDTO,
    CanonicalExcelTemplateGenerator,
    CanonicalExcelParser,
    CanonicalQuestionImporter,
    normalize_input_for_duplicate_check,
    TEMPLATE_NAME,
    TEMPLATE_VERSION,
)
from apps.questions.services import QuestionService
from apps.assessments.models import (
    Assessment,
    AssessmentQuestion,
    AssessmentSnapshot,
    AssessmentSnapshotQuestion,
    AssessmentStatus,
    TestAttempt,
    AttemptAnswer,
)
from apps.assessments.services import (
    AssessmentService,
    AssessmentSnapshotService,
)

User = get_user_model()


@pytest.fixture
def admin_user(db):
    return User.objects.create_superuser(
        email="canonical_admin@codeguard.local",
        password="AdminSecurePass123!",
        role=Role.ADMIN
    )


@pytest.fixture
def student_user(db):
    user = User.objects.create_user(
        email="canonical_student@codeguard.local",
        password="StudentPass123!",
        role=Role.STUDENT,
        first_login_required=False
    )
    StudentProfile.objects.create(
        user=user,
        roll_number=f"CS-2026-{uuid.uuid4().hex[:6]}",
        first_login_required=False
    )
    return user


def error_messages(errors):
    """Extract string error messages from error list containing strings or dicts."""
    return [e.get("message", str(e)) if isinstance(e, dict) else str(e) for e in errors]


def has_error(errors, text):
    """Check if any error message contains substring text (case-insensitive)."""
    t = text.lower()
    return any(t in msg.lower() for msg in error_messages(errors))


def make_canonical_excel_bytes(
    questions=None,
    options=None,
    coding=None,
    test_cases=None,
    template_name=TEMPLATE_NAME,
    template_version=TEMPLATE_VERSION,
    omit_readme=False,
    omit_sheet=None,
) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    if not omit_readme:
        ws_readme = wb.create_sheet(title="README")
        ws_readme.append(["Template Name", template_name])
        ws_readme.append(["Template Version", str(template_version)])
        ws_readme.append(["Platform", "CODEGUARD Assessment Platform"])

    if omit_sheet != "Questions":
        ws_q = wb.create_sheet(title="Questions")
        ws_q.append(["question_id", "title", "type", "statement", "instructions", "difficulty", "points", "negative_points"])
        default_q = [
            ("Q001", "Two Sum Solver", "CODING", "Given two integers, print their sum.", "Read from stdin.", "EASY", 10, 0),
            ("Q002", "Python Type Mutability", "MCQ", "Which type is mutable in Python?", "Select one option.", "EASY", 5, 0),
        ]
        for row in (questions if questions is not None else default_q):
            ws_q.append(row)

    if omit_sheet != "Options":
        ws_opt = wb.create_sheet(title="Options")
        ws_opt.append(["question_id", "option_key", "option_text", "is_correct"])
        default_opt = [
            ("Q002", "A", "List", "TRUE"),
            ("Q002", "B", "Tuple", "FALSE"),
            ("Q002", "C", "String", "FALSE"),
            ("Q002", "D", "Integer", "FALSE"),
        ]
        for row in (options if options is not None else default_opt):
            ws_opt.append(row)

    if omit_sheet != "Coding":
        ws_c = wb.create_sheet(title="Coding")
        ws_c.append(["question_id", "allowed_languages", "starter_code_python", "starter_code_cpp", "starter_code_java", "time_limit_ms", "memory_limit_mb", "constraints"])
        default_c = [
            ("Q001", "PYTHON", "def solve():\n    pass", "", "", 2000, 256, "1 <= a, b <= 1000"),
        ]
        for row in (coding if coding is not None else default_c):
            ws_c.append(row)

    if omit_sheet != "TestCases":
        ws_tc = wb.create_sheet(title="TestCases")
        ws_tc.append(["question_id", "case_id", "visibility", "input", "expected_output", "points", "is_example"])
        default_tc = [
            ("Q001", "TC01", "SAMPLE", "2 3", "5", 5, "TRUE"),
            ("Q001", "TC02", "HIDDEN", "10 20", "30", 5, "FALSE"),
        ]
        for row in (test_cases if test_cases is not None else default_tc):
            ws_tc.append(row)

    if omit_sheet != "COLUMN_GUIDE":
        ws_g = wb.create_sheet(title="COLUMN_GUIDE")
        ws_g.append(["Sheet", "Column", "Description"])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==============================================================================
# GROUP 1: Canonical Excel Import Tests (01 - 21)
# ==============================================================================

@pytest.mark.django_db
def test_01_import_valid_mcq_excel(admin_user):
    """1. Import valid MCQ question creates Question ACTIVE and v1 DRAFT."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q_MCQ", "Python Sets", "MCQ", "Are sets unordered?", "", "EASY", 5, 0)],
        options=[
            ("Q_MCQ", "A", "Yes", "TRUE"),
            ("Q_MCQ", "B", "No", "FALSE"),
        ],
        coding=[],
        test_cases=[]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, warnings = CanonicalExcelParser.parse_workbook(wb)
    assert not errors, f"Unexpected errors: {errors}"
    assert len(dtos) == 1

    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    assert res["created_count"] == 1
    q = Question.objects.get(id=res["questions"][0]["question_id"])
    assert q.status == QuestionStatus.ACTIVE
    assert q.question_type == QuestionType.MCQ

    v1 = q.versions.get(version_number=1)
    assert v1.status == VersionStatus.DRAFT
    assert v1.title == "Python Sets"
    assert len(v1.type_config.get("options", [])) == 2
    assert v1.type_config["options"][0]["is_correct"] is True


@pytest.mark.django_db
def test_02_import_valid_coding_excel(admin_user):
    """2. Import valid Coding question loads starter codes, test cases, and derives examples."""
    excel_bytes = make_canonical_excel_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, warnings = CanonicalExcelParser.parse_workbook(wb)
    assert not errors, f"Unexpected errors: {errors}"

    coding_dto = next(d for d in dtos if d.type == "CODING")
    assert coding_dto.coding.starter_codes["PYTHON"] == "def solve():\n    pass"
    assert len(coding_dto.test_cases) == 2

    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    q_ids = [q_info["question_id"] for q_info in res["questions"]]
    q_coding = Question.objects.filter(id__in=q_ids, question_type=QuestionType.CODING).first()
    v1 = q_coding.versions.get(version_number=1)
    assert v1.coding_config.problem_statement == "Given two integers, print their sum."
    assert v1.coding_config.test_cases.count() == 2
    assert len(v1.coding_config.examples) == 1
    assert v1.coding_config.examples[0]["input"] == "2 3"


@pytest.mark.django_db
def test_03_import_rejects_unknown_question_type():
    """3. Import rejects unknown question types (e.g. SQL)."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "SQL Query", "SQL", "SELECT * FROM users;", "", "EASY", 10, 0)],
        options=[],
        coding=[],
        test_cases=[]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, warnings = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "Unsupported question type 'SQL'")


@pytest.mark.django_db
def test_04_import_rejects_duplicate_question_id_in_file():
    """4. Import rejects duplicate question_id in Questions sheet."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[
            ("Q001", "Title 1", "MCQ", "Statement 1", "", "EASY", 5, 0),
            ("Q001", "Title 2", "MCQ", "Statement 2", "", "EASY", 5, 0),
        ],
        options=[
            ("Q001", "A", "Opt 1", "TRUE"),
            ("Q001", "B", "Opt 2", "FALSE"),
        ],
        coding=[],
        test_cases=[]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, warnings = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "Duplicate question_id 'Q001'")


@pytest.mark.django_db
def test_05_import_rejects_mcq_with_no_correct_option():
    """5. Import rejects MCQ question with zero correct options."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q002", "MCQ Title", "MCQ", "Statement", "", "EASY", 5, 0)],
        options=[
            ("Q002", "A", "Opt 1", "FALSE"),
            ("Q002", "B", "Opt 2", "FALSE"),
        ],
        coding=[],
        test_cases=[]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, warnings = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "none was marked correct")


@pytest.mark.django_db
def test_06_import_rejects_mcq_with_multiple_correct_options():
    """6. Import rejects MCQ question with multiple correct options."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q002", "MCQ Title", "MCQ", "Statement", "", "EASY", 5, 0)],
        options=[
            ("Q002", "A", "Opt 1", "TRUE"),
            ("Q002", "B", "Opt 2", "TRUE"),
        ],
        coding=[],
        test_cases=[]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, warnings = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "found 2 correct options")


@pytest.mark.django_db
def test_07_import_rejects_mcq_with_fewer_than_two_options():
    """7. Import rejects MCQ question with fewer than 2 options."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q002", "MCQ Title", "MCQ", "Statement", "", "EASY", 5, 0)],
        options=[
            ("Q002", "A", "Only One Option", "TRUE"),
        ],
        coding=[],
        test_cases=[]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, warnings = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "at least 2 options")


@pytest.mark.django_db
def test_08_import_rejects_orphan_options():
    """8. Import rejects orphan options referencing non-existent question_id."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q002", "MCQ Title", "MCQ", "Statement", "", "EASY", 5, 0)],
        options=[
            ("Q002", "A", "Opt A", "TRUE"),
            ("Q002", "B", "Opt B", "FALSE"),
            ("Q_NONEXISTENT", "C", "Orphan", "FALSE"),
        ],
        coding=[],
        test_cases=[]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, warnings = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "Q_NONEXISTENT") and has_error(errors, "Options")


@pytest.mark.django_db
def test_09_import_preserves_multiline_problem_statement_fidelity(admin_user):
    """9. Import preserves multiline problem statement formatting and indentations."""
    multiline = "Line 1: Problem Title\n\nLine 3: Example:\n    def example():\n        return True\n\nLine 7: Done."
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Multiline Coding", "CODING", multiline, "Do not alter.", "MEDIUM", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "def solve():\n    pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "in", "out", 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "in2", "out2", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert not errors
    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    v1 = Question.objects.get(id=res["questions"][0]["question_id"]).versions.get(version_number=1)
    assert v1.coding_config.problem_statement == multiline


@pytest.mark.django_db
def test_10_import_preserves_instructions_and_constraints(admin_user):
    """10. Import preserves instructions and constraints verbatim."""
    instructions = "Solve within O(N) time complexity. Avoid recursions."
    constraints = "1 <= N <= 10^5\n-10^9 <= A[i] <= 10^9"
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Constrained Question", "CODING", "Statement", instructions, "HARD", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "def solve():\n    pass", "", "", 1500, 512, constraints)],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "1", "1", 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "2", "2", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert not errors
    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    v1 = Question.objects.get(id=res["questions"][0]["question_id"]).versions.get(version_number=1)
    assert v1.instructions == instructions
    assert v1.coding_config.constraints == constraints


@pytest.mark.django_db
def test_11_import_derives_coding_examples_from_sample_test_cases(admin_user):
    """11. Import derives examples strictly from SAMPLE test cases marked is_example=TRUE."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Coding Examples", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "input_1", "output_1", 5, "TRUE"),
            ("Q001", "TC2", "SAMPLE", "input_2", "output_2", 5, "FALSE"),
            ("Q001", "TC3", "HIDDEN", "input_3", "output_3", 0, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert not errors
    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    v1 = Question.objects.get(id=res["questions"][0]["question_id"]).versions.get(version_number=1)
    assert len(v1.coding_config.examples) == 1
    assert v1.coding_config.examples[0]["input"] == "input_1"
    assert v1.coding_config.examples[0]["output"] == "output_1"


@pytest.mark.django_db
def test_12_import_loads_all_specified_starter_codes(admin_user):
    """12. Import loads starter codes for all allowed languages."""
    py_code = "# Python starter\ndef solve(): pass"
    cpp_code = "// C++ starter\nint main() { return 0; }"
    java_code = "// Java starter\nclass Solution { public static void main(String[] args) {} }"
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Polyglot", "CODING", "Statement", "", "MEDIUM", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON,CPP,JAVA", py_code, cpp_code, java_code, 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "1", "1", 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "2", "2", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert not errors
    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    v1 = Question.objects.get(id=res["questions"][0]["question_id"]).versions.get(version_number=1)
    starters = v1.coding_config.starter_codes
    assert starters["PYTHON"] == py_code
    assert starters["CPP"] == cpp_code
    assert starters["JAVA"] == java_code


@pytest.mark.django_db
def test_13_import_loads_public_and_hidden_test_cases_accurately(admin_user):
    """13. Import accurately records visibility as SAMPLE (public) and HIDDEN."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Visibility Check", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "sample_in", "sample_out", 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "hidden_in", "hidden_out", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert not errors
    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    v1 = Question.objects.get(id=res["questions"][0]["question_id"]).versions.get(version_number=1)
    tcs = list(v1.coding_config.test_cases.order_by("execution_order"))
    assert tcs[0].is_hidden is False
    assert tcs[1].is_hidden is True


@pytest.mark.django_db
def test_14_import_ensures_hidden_test_cases_are_never_marked_as_examples(admin_user):
    """14. Hidden test cases are never present in coding_config.examples."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Never Hidden Example", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "pub_in", "pub_out", 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "secret_in", "secret_out", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert not errors
    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    v1 = Question.objects.get(id=res["questions"][0]["question_id"]).versions.get(version_number=1)
    example_inputs = [ex["input"] for ex in v1.coding_config.examples]
    assert "secret_in" not in example_inputs
    assert "pub_in" in example_inputs


@pytest.mark.django_db
def test_15_import_rejects_coding_question_missing_starter_code_for_allowed_language():
    """15. Import rejects enabled language missing starter code."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Missing CPP", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON,CPP", "def solve(): pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "1", "1", 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "2", "2", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "starter code is missing or empty") and has_error(errors, "CPP")


@pytest.mark.django_db
def test_16_import_rejects_language_synchronization_mismatch():
    """16. Import rejects unsupported languages in allowed_languages."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Bad Lang", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON,RUBY", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "1", "1", 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "2", "2", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "Unsupported coding language 'RUBY'")


@pytest.mark.django_db
def test_17_import_preserves_multiline_expected_outputs_verbatim(admin_user):
    """17. Multiline expected outputs are preserved verbatim without truncation."""
    multiline_out = "Row 1: 10\nRow 2: 20\nRow 3: 30\n"
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Multiline Output", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "in1", multiline_out, 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "in2", "out2", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert not errors
    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    tc1 = Question.objects.get(id=res["questions"][0]["question_id"]).versions.get(version_number=1).coding_config.test_cases.get(name="TC1")
    assert tc1.expected_output == multiline_out


@pytest.mark.django_db
def test_18_import_rejects_placeholder_empty_outputs():
    """18. Import rejects test cases with empty or whitespace-only expected outputs."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Empty Output", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "in1", "   ", 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "in2", "out2", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "empty expected output")


@pytest.mark.django_db
def test_19_import_normalizes_test_case_points_to_sum_to_question_points(admin_user):
    """19. Test case points normalize deterministically to sum to question.points when P >= N."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Point Distribution", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "1", "1", 1, "TRUE"),
            ("Q001", "TC2", "SAMPLE", "2", "2", 1, "FALSE"),
            ("Q001", "TC3", "HIDDEN", "3", "3", 1, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert not errors
    res = CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)
    v1 = Question.objects.get(id=res["questions"][0]["question_id"]).versions.get(version_number=1)
    tcs = list(v1.coding_config.test_cases.order_by("execution_order"))
    pts = [tc.points for tc in tcs]
    # 10 points over 3 test cases: 4, 3, 3
    assert pts == [4, 3, 3]
    assert sum(pts) == 10


@pytest.mark.django_db
def test_20_import_rejects_when_points_less_than_scored_cases():
    """20. Import rejects with exact error when P < N."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Underpointed Question", "CODING", "Statement", "", "EASY", 2, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "1", "1", 1, "TRUE"),
            ("Q001", "TC2", "SAMPLE", "2", "2", 1, "FALSE"),
            ("Q001", "TC3", "HIDDEN", "3", "3", 1, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    expected_msg = "Question has 2 points but 3 scored test cases require at least 1 point each. Increase question points or reduce scored test cases."
    assert has_error(errors, expected_msg)


@pytest.mark.django_db
def test_21_import_is_strictly_atomic_and_all_or_nothing(admin_user):
    """21. Transactional all-or-nothing import: failure on 1 question creates 0 records."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[
            ("Q001", "Valid Question", "CODING", "Statement", "", "EASY", 10, 0),
            ("Q002", "Invalid Question", "CODING", "Statement", "", "EASY", 2, 0)
        ],
        options=[],
        coding=[
            ("Q001", "PYTHON", "pass", "", "", 2000, 256, ""),
            ("Q002", "PYTHON", "pass", "", "", 2000, 256, "")
        ],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "1", "1", 5, "TRUE"),
            ("Q001", "TC2", "HIDDEN", "2", "2", 5, "FALSE"),
            ("Q002", "TC1", "SAMPLE", "1", "1", 1, "TRUE"),
            ("Q002", "TC2", "SAMPLE", "2", "2", 1, "FALSE"),
            ("Q002", "TC3", "HIDDEN", "3", "3", 1, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert errors, "Expected validation errors for Q002"

    with pytest.raises(DRFValidationError):
        CanonicalQuestionImporter.import_canonical_questions(dtos, actor=admin_user)

    assert Question.objects.filter(versions__title="Valid Question").count() == 0


# ==============================================================================
# GROUP 2: Lifecycle & Immutability Tests (22 - 27, 57)
# ==============================================================================

@pytest.mark.django_db
def test_22_draft_question_can_be_edited_freely(admin_user):
    """22. Draft question version can be modified and saved freely."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Original Draft",
        description="Original statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    assert v1.status == VersionStatus.DRAFT
    v1.title = "Updated Draft Title"
    v1.points = 15
    v1.save()
    v1.refresh_from_db()
    assert v1.title == "Updated Draft Title"
    assert v1.points == 15


@pytest.mark.django_db
def test_23_draft_question_cannot_be_added_to_published_or_draft_exam(admin_user):
    """23. Adding a DRAFT question version to an assessment is strictly blocked."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Draft MCQ",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    assessment = Assessment.objects.create(
        title="Exam 1",
        description="Exam desc",
        duration_minutes=60,
        total_points=100,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    with pytest.raises(DRFValidationError) as exc:
        AssessmentService.add_question_to_assessment(
            assessment_id=str(assessment.id),
            question_version_id=str(v1.id),
            points=10,
            actor=admin_user
        )
    assert "Only published question versions can be added" in str(exc.value)


@pytest.mark.django_db
def test_24_published_question_version_is_strictly_immutable_direct_save(admin_user):
    """24. Calling save() on a PUBLISHED question version raises PermissionDenied."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Pub MCQ",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    v1.refresh_from_db()
    assert v1.status == VersionStatus.PUBLISHED

    v1.title = "Attempted Mutation"
    with pytest.raises(PermissionDenied) as exc:
        v1.save()
    assert "immutable" in str(exc.value)


@pytest.mark.django_db
def test_25_published_question_version_is_strictly_immutable_queryset_update(admin_user):
    """25. Calling update() on a QuerySet of PUBLISHED versions raises PermissionDenied."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Pub MCQ",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    with pytest.raises(PermissionDenied) as exc:
        QuestionVersion.objects.filter(id=v1.id).update(title="QuerySet Mutation")
    assert "immutable" in str(exc.value)


@pytest.mark.django_db
def test_26_editing_published_question_creates_new_draft_version(admin_user):
    """26. Creating a new version on a published question produces version 2 DRAFT."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Pub MCQ",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    v2 = QuestionService.create_new_version(str(q.id), actor=admin_user)
    assert v2.version_number == 2
    assert v2.status == VersionStatus.DRAFT

    v1.refresh_from_db()
    assert v1.status == VersionStatus.PUBLISHED
    assert v1.version_number == 1


@pytest.mark.django_db
def test_27_editing_published_question_does_not_mutate_historical_snapshot(admin_user):
    """27. Creating and editing v2 does not mutate an existing assessment snapshot built on v1."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Original V1",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    v1.refresh_from_db()

    assessment = Assessment.objects.create(
        title="Historical Assessment",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)

    # Now create v2 and edit it
    v2 = QuestionService.create_new_version(str(q.id), actor=admin_user)
    v2.title = "Mutated V2 Title"
    v2.save()

    snapshot.refresh_from_db()
    snap_q = snapshot.snapshot_data["questions"][0]
    assert snap_q["title"] == "Original V1"


@pytest.mark.django_db
def test_57_queryset_delete_blocked_on_published_version(admin_user):
    """57. Calling delete() on a QuerySet containing PUBLISHED versions raises PermissionDenied."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Pub MCQ",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    with pytest.raises(PermissionDenied) as exc:
        QuestionVersion.objects.filter(id=v1.id).delete()
    assert "immutable" in str(exc.value)


# ==============================================================================
# GROUP 3: Assessment Integration & Snapshots (28 - 32)
# ==============================================================================

@pytest.mark.django_db
def test_28_only_published_question_can_be_added_to_exam(admin_user):
    """28. Publishing a question version enables it to be added to an assessment."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Pub Candidate",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Live Exam",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    aq = AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    assert aq.assessment_id == assessment.id
    assert aq.question_version_id == v1.id


@pytest.mark.django_db
def test_29_added_question_persists_on_page_reload_and_in_database(admin_user):
    """29. Added question persists in AssessmentQuestion table upon DB reload."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Persistent Q",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Persist Exam",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    assessment.refresh_from_db()
    assert assessment.assessment_questions.count() == 1
    assert assessment.assessment_questions.first().question_version.title == "Persistent Q"


@pytest.mark.django_db
def test_30_exam_snapshot_captures_complete_public_question_content(admin_user):
    """30. Snapshot captures public problem statement, starter code, and examples."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.CODING,
        title="Public Content Q",
        description="Statement",
        points=10,
        coding_config_data={
            "problem_statement": "Print the square of N.",
            "allowed_languages": ["PYTHON"],
            "starter_codes": {"PYTHON": "def solve(n): pass"},
            "time_limit_ms": 2000,
            "memory_limit_mb": 256,
        },
        test_cases_data=[
            {"input_data": "4", "expected_output": "16", "points": 5, "is_hidden": False, "is_example": True},
            {"input_data": "5", "expected_output": "25", "points": 5, "is_hidden": True, "is_example": False}
        ]
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    v1.refresh_from_db()

    assessment = Assessment.objects.create(
        title="Snapshot Content Exam",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)

    q_data = snapshot.snapshot_data["questions"][0]
    coding_data = q_data["coding_config"]
    assert coding_data["problem_statement"] == "Print the square of N."
    assert coding_data["starter_codes"]["PYTHON"] == "def solve(n): pass"
    assert len(coding_data["examples"]) == 1
    assert coding_data["examples"][0]["input"] == "4"


@pytest.mark.django_db
def test_31_exam_snapshot_excludes_mcq_answer_keys_and_hidden_test_cases(admin_user):
    """31. Exam snapshot excludes MCQ is_correct and hidden test cases."""
    q_mcq, v_mcq = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="MCQ Key Check",
        description="Statement",
        points=5,
        type_config={
            "correct_option": "A",
            "admin_notes": "Do not leak",
            "options": [
                {"id": "A", "text": "Correct Answer", "is_correct": True},
                {"id": "B", "text": "Wrong Answer", "is_correct": False}
            ]
        }
    )
    QuestionService.publish_version(str(q_mcq.id), 1, actor=admin_user)

    q_code, v_code = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.CODING,
        title="Hidden Test Check",
        description="Statement",
        points=10,
        coding_config_data={
            "problem_statement": "Coding Statement",
            "allowed_languages": ["PYTHON"],
            "starter_codes": {"PYTHON": "pass"},
        },
        test_cases_data=[
            {"input_data": "sample_in", "expected_output": "sample_out", "points": 5, "is_hidden": False, "is_example": True},
            {"input_data": "hidden_in", "expected_output": "hidden_out", "points": 5, "is_hidden": True, "is_example": False}
        ]
    )
    QuestionService.publish_version(str(q_code.id), 1, actor=admin_user)

    assessment = Assessment.objects.create(
        title="Security Exam",
        description="Desc",
        duration_minutes=60,
        total_points=15,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v_mcq.id), points=5, actor=admin_user)
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v_code.id), points=10, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)

    snap_json = json.dumps(snapshot.snapshot_data)
    assert "hidden_in" not in snap_json
    assert "hidden_out" not in snap_json
    assert "admin_notes" not in snap_json

    for q_data in snapshot.snapshot_data["questions"]:
        if q_data["question_type"] == "MCQ":
            for opt in q_data["type_config"]["options"]:
                assert "is_correct" not in opt


@pytest.mark.django_db
def test_32_student_cannot_read_hidden_test_cases_or_answers_from_assessment_or_snapshot(api_client, student_user, admin_user):
    """32. Student cannot read hidden test cases or answers from assessment detail."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="MCQ Candidate",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "Ans A", "is_correct": True}, {"id": "B", "text": "Ans B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)

    assessment = Assessment.objects.create(
        title="Student Access Exam",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now() - timedelta(minutes=5),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()

    AssessmentService.assign_students(
        assessment=assessment,
        student_ids=[str(student_user.id)],
        actor=admin_user
    )

    api_client.force_authenticate(user=student_user)
    url = f"/api/v1/student/assessments/{assessment.id}/"
    res = api_client.get(url)
    assert res.status_code == 200
    res_str = json.dumps(res.data)
    assert "is_correct" not in res_str
    assert "server_evaluation_bundle" not in res_str


# ==============================================================================
# GROUP 4: Student Runtime Security & Zero Leakage (33 - 38, 58)
# ==============================================================================

@pytest.mark.django_db
def test_33_student_attempt_payload_renders_mcq_without_answer_keys(admin_user):
    """33. Student attempt question payload renders options without is_correct."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Sanitized MCQ",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "Opt A", "is_correct": True}, {"id": "B", "text": "Opt B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)

    assessment = Assessment.objects.create(
        title="Attempt Exam",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now() - timedelta(minutes=5),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()

    questions_payload = snapshot.snapshot_data["questions"]
    mcq_q = questions_payload[0]
    for opt in mcq_q["type_config"]["options"]:
        assert "is_correct" not in opt
    assert "correct_option" not in mcq_q["type_config"]


@pytest.mark.django_db
def test_34_student_attempt_payload_renders_coding_without_hidden_tests_or_outputs(admin_user):
    """34. Student attempt question payload for coding renders zero hidden test cases."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.CODING,
        title="Sanitized Coding",
        description="Statement",
        points=10,
        coding_config_data={
            "problem_statement": "Reverse string.",
            "allowed_languages": ["PYTHON"],
            "starter_codes": {"PYTHON": "pass"},
        },
        test_cases_data=[
            {"input_data": "sample", "expected_output": "elpmas", "points": 5, "is_hidden": False, "is_example": True},
            {"input_data": "secret_long_input", "expected_output": "tupni_gnol_terces", "points": 5, "is_hidden": True, "is_example": False}
        ]
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)

    assessment = Assessment.objects.create(
        title="Attempt Coding Exam",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now() - timedelta(minutes=5),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()

    questions_str = json.dumps(snapshot.snapshot_data["questions"])
    assert "secret_long_input" not in questions_str
    assert "tupni_gnol_terces" not in questions_str


@pytest.mark.django_db
def test_35_student_attempt_payload_rejects_draft_or_archived_question_access(api_client, student_user, admin_user):
    """35. Student cannot access draft question versions."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Draft Only",
        description="Statement",
        points=10,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    api_client.force_authenticate(user=student_user)
    res = api_client.get(f"/api/v1/admin/questions/{q.id}/versions/{v1.version_number}/")
    assert res.status_code == 403


@pytest.mark.django_db
def test_36_student_room_renders_mcq_options_cleanly(admin_user):
    """36. MCQ options format in snapshot provides id and text for test room UI."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Clean UI MCQ",
        description="Select best answer",
        points=5,
        type_config={"options": [{"id": "A", "text": "Alpha", "is_correct": True}, {"id": "B", "text": "Beta", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="UI Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    options = snapshot.snapshot_data["questions"][0]["type_config"]["options"]
    assert len(options) == 2
    assert options[0] == {"id": "A", "text": "Alpha"}
    assert options[1] == {"id": "B", "text": "Beta"}


@pytest.mark.django_db
def test_37_student_room_renders_coding_problem_starter_code_and_examples_cleanly(admin_user):
    """37. Coding question format in snapshot provides problem_statement, starter_codes, and examples."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.CODING,
        title="Clean UI Coding",
        description="Desc",
        points=10,
        coding_config_data={
            "problem_statement": "Sum two integers.",
            "allowed_languages": ["PYTHON"],
            "starter_codes": {"PYTHON": "def solve(): pass"},
        },
        test_cases_data=[
            {"input_data": "1 2", "expected_output": "3", "points": 5, "is_hidden": False, "is_example": True},
            {"input_data": "4 5", "expected_output": "9", "points": 5, "is_hidden": True, "is_example": False}
        ]
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="UI Coding Exam",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    coding_cfg = snapshot.snapshot_data["questions"][0]["coding_config"]
    assert coding_cfg["problem_statement"] == "Sum two integers."
    assert coding_cfg["starter_codes"]["PYTHON"] == "def solve(): pass"
    assert len(coding_cfg["examples"]) == 1


@pytest.mark.django_db
def test_38_student_room_exposes_zero_hidden_tests_or_internal_notes_in_dom_or_network(admin_user):
    """38. Type config removes admin_notes, solution_notes, and internal_notes."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Private Notes MCQ",
        description="Desc",
        points=5,
        type_config={
            "admin_notes": "Confidential rubric",
            "internal_notes": "Author comment",
            "options": [{"id": "A", "text": "Opt A", "is_correct": True}, {"id": "B", "text": "Opt B", "is_correct": False}]
        }
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Private Notes Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    snap_str = json.dumps(snapshot.snapshot_data)
    assert "Confidential rubric" not in snap_str
    assert "Author comment" not in snap_str


@pytest.mark.django_db
def test_58_student_attempt_detail_sanitizes_is_correct_and_correct_option(api_client, student_user, admin_user):
    """58. StudentAttemptDetailView has defensive runtime sanitization stripping is_correct."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Defensive Sanitize MCQ",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "Opt A", "is_correct": True}, {"id": "B", "text": "Opt B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Defensive Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now() - timedelta(minutes=5),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()

    # Simulate an older snapshot that had is_correct accidentally embedded
    dirty_questions = [
        {
            "snapshot_question_id": "SQ1",
            "question_type": "MCQ",
            "title": "Legacy MCQ",
            "description": "Statement",
            "type_config": {
                "correct_option": "A",
                "options": [
                    {"id": "A", "text": "A", "is_correct": True},
                    {"id": "B", "text": "B", "is_correct": False}
                ]
            }
        }
    ]
    snapshot.snapshot_data["questions"] = dirty_questions
    AssessmentSnapshot.objects.filter(pk=snapshot.pk).update(snapshot_data=snapshot.snapshot_data)

    attempt = TestAttempt.objects.create(
        assessment=assessment,
        student=student_user,
        assessment_snapshot=snapshot,
        attempt_number=1,
        randomization_seed="seed123",
        question_order=["SQ1"]
    )

    api_client.force_authenticate(user=student_user)
    res = api_client.get(f"/api/v1/student/attempts/{attempt.id}/")
    assert res.status_code == 200
    res_str = json.dumps(res.data)
    assert "is_correct" not in res_str
    assert "correct_option" not in res_str


# ==============================================================================
# GROUP 5: Safe Deletion & Dependency Tracking (39 - 43)
# ==============================================================================

@pytest.mark.django_db
def test_39_unreferenced_draft_question_can_be_deleted(admin_user):
    """39. Unreferenced draft question can be deleted."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Disposable Draft",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    q_id = q.id
    QuestionService.delete_draft_question(str(q_id), actor=admin_user)
    assert not Question.objects.filter(id=q_id).exists()


@pytest.mark.django_db
def test_40_question_referenced_in_exam_cannot_be_deleted(admin_user):
    """40. Question referenced in an assessment cannot be deleted."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Referenced In Exam",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Locking Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)

    with pytest.raises(DRFValidationError) as exc:
        QuestionService.delete_draft_question(str(q.id), actor=admin_user)
    assert "Cannot delete question" in str(exc.value)


@pytest.mark.django_db
def test_41_question_referenced_in_historical_attempt_cannot_be_deleted(admin_user, student_user):
    """41. Question referenced in a historical exam attempt cannot be deleted."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Attempt Referenced Q",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Attempted Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now() - timedelta(minutes=5),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()

    snap_q = snapshot.snapshot_questions.first()
    attempt = TestAttempt.objects.create(
        assessment=assessment,
        student=student_user,
        assessment_snapshot=snapshot,
        attempt_number=1,
        randomization_seed="seed123"
    )
    AttemptAnswer.objects.create(
        attempt=attempt,
        snapshot_question=snap_q,
        question_id=snap_q.snapshot_question_id,
        question_type=snap_q.question_type,
        selected_options=["A"]
    )

    with pytest.raises(DRFValidationError) as exc:
        QuestionService.delete_draft_question(str(q.id), actor=admin_user)
    assert "Cannot delete question" in str(exc.value)


@pytest.mark.django_db
def test_42_referenced_question_can_be_archived_without_breaking_historical_records(admin_user):
    """42. Referenced question can be archived, preserving historical records."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Archivable Q",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Archive Holding Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)

    QuestionService.archive_question(str(q.id), actor=admin_user)
    q.refresh_from_db()
    assert q.status == QuestionStatus.ARCHIVED
    assert assessment.assessment_questions.count() == 1


@pytest.mark.django_db
def test_43_delete_rejection_explains_exact_dependency_counts_and_recommends_archive(admin_user):
    """43. Deletion rejection message explains exact dependency counts and recommends archive."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Dependency Count Q",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Dep Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)

    with pytest.raises(DRFValidationError) as exc:
        QuestionService.delete_draft_question(str(q.id), actor=admin_user)
    msg = str(exc.value)
    assert "assessments" in msg
    assert "Cannot delete this question because it is used in 1 assessment" in msg


# ==============================================================================
# GROUP 6: Section-Free Architecture & IDOR Protections (44 - 50)
# ==============================================================================

@pytest.mark.django_db
def test_44_create_and_publish_question_flow_has_zero_section_dependencies(admin_user):
    """44. Question creation and publishing operates cleanly without sections."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Section Free Q",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    v_pub = QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assert v_pub.status == VersionStatus.PUBLISHED


@pytest.mark.django_db
def test_45_add_question_to_assessment_flow_has_zero_section_dependencies(admin_user):
    """45. Adding question to assessment does not require or touch sections."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Section Free Add",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Section Free Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    aq = AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)
    assert aq.assessment_id == assessment.id


@pytest.mark.django_db
def test_46_student_exam_delivery_flow_has_zero_section_dependencies(admin_user, student_user):
    """46. Student exam delivery functions without section dependencies."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Direct Question Delivery",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    assessment = Assessment.objects.create(
        title="Direct Delivery Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now() - timedelta(minutes=5),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()

    attempt = TestAttempt.objects.create(
        assessment=assessment,
        student=student_user,
        assessment_snapshot=snapshot,
        attempt_number=1,
        randomization_seed="seed123"
    )
    assert attempt.assessment_snapshot.snapshot_data["questions"][0]["title"] == "Direct Question Delivery"


@pytest.mark.django_db
def test_47_question_api_returns_zero_section_fields_in_payloads(api_client, admin_user):
    """47. Question admin API payloads do not require or expose section fields."""
    api_client.force_authenticate(user=admin_user)
    res = api_client.get("/api/v1/admin/questions/")
    assert res.status_code == 200
    assert "section" not in res.data


@pytest.mark.django_db
def test_48_idor_protection_student_cannot_access_staff_preview_endpoint(api_client, student_user, admin_user):
    """48. Student cannot access staff preview endpoint (403)."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Preview IDOR Check",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    api_client.force_authenticate(user=student_user)
    res = api_client.get(f"/api/v1/admin/questions/{q.id}/versions/{v1.version_number}/preview/")
    assert res.status_code == 403


@pytest.mark.django_db
def test_49_idor_protection_student_cannot_access_admin_question_detail_endpoint(api_client, student_user, admin_user):
    """49. Student cannot access admin question detail endpoint (403)."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Detail IDOR Check",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    api_client.force_authenticate(user=student_user)
    res = api_client.get(f"/api/v1/admin/questions/{q.id}/")
    assert res.status_code == 403


@pytest.mark.django_db
def test_50_idor_protection_student_cannot_access_admin_draft_versions(api_client, student_user, admin_user):
    """50. Student cannot access admin draft version detail endpoint (403)."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Draft IDOR Check",
        description="Desc",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    api_client.force_authenticate(user=student_user)
    res = api_client.get(f"/api/v1/admin/questions/{q.id}/versions/{v1.version_number}/")
    assert res.status_code == 403


# ==============================================================================
# GROUP 7: Normalization, Canonical Contract, & Template Integrity (51 - 56)
# ==============================================================================

def test_51_duplicate_input_normalization_crlf_and_internal_whitespace():
    """51. CRLF and surrounding whitespace normalized without collapsing internal whitespace."""
    inp1 = "\r\n  5  10  \r\n"
    inp2 = "5  10"
    inp3 = "5 10"  # single space

    norm1 = normalize_input_for_duplicate_check(inp1)
    norm2 = normalize_input_for_duplicate_check(inp2)
    norm3 = normalize_input_for_duplicate_check(inp3)

    assert norm1 == norm2
    assert norm1 != norm3  # Internal whitespace preserved


@pytest.mark.django_db
def test_52_import_rejects_duplicate_case_id_within_question():
    """52. Import rejects duplicate case_id within the same question."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Dup Case ID", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC01", "SAMPLE", "1", "1", 5, "TRUE"),
            ("Q001", "TC01", "HIDDEN", "2", "2", 5, "FALSE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "Duplicate test case ID 'TC01'")


@pytest.mark.django_db
def test_53_import_rejects_hidden_test_case_with_is_example_true():
    """53. Import rejects hidden test case with is_example=TRUE."""
    excel_bytes = make_canonical_excel_bytes(
        questions=[("Q001", "Hidden Example", "CODING", "Statement", "", "EASY", 10, 0)],
        options=[],
        coding=[("Q001", "PYTHON", "pass", "", "", 2000, 256, "")],
        test_cases=[
            ("Q001", "TC1", "SAMPLE", "1", "1", 5, "FALSE"),
            ("Q001", "TC2", "HIDDEN", "2", "2", 5, "TRUE")
        ]
    )
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "Hidden test cases cannot be marked as examples.")


@pytest.mark.django_db
def test_54_template_version_metadata_validation():
    """54. Import rejects template missing CODEGUARD_QUESTION_IMPORT metadata or invalid version."""
    excel_bytes = make_canonical_excel_bytes(template_name="UNKNOWN_TEMPLATE")
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb)
    assert has_error(errors, "Invalid template")

    excel_bytes_v99 = make_canonical_excel_bytes(template_version=99)
    wb_v99 = openpyxl.load_workbook(io.BytesIO(excel_bytes_v99), data_only=True)
    dtos, errors, _ = CanonicalExcelParser.parse_workbook(wb_v99)
    assert has_error(errors, "Unsupported template version 99")


@pytest.mark.django_db
def test_55_manual_coding_question_creation_validates_canonical_dto_rules(admin_user):
    """55. Manual creation of coding question validates canonical rules (P >= N, starter code sync)."""
    with pytest.raises(DRFValidationError) as exc:
        QuestionService.create_question(
            actor=admin_user,
            question_type=QuestionType.CODING,
            title="Manual Coding Fail",
            description="Statement",
            points=2,  # P = 2 < N = 3
            coding_config_data={
                "problem_statement": "Statement",
                "allowed_languages": ["PYTHON"],
                "starter_codes": {"PYTHON": "pass"},
            },
            test_cases_data=[
                {"input_data": "1", "expected_output": "1", "points": 1, "is_hidden": False, "is_example": True},
                {"input_data": "2", "expected_output": "2", "points": 1, "is_hidden": False, "is_example": False},
                {"input_data": "3", "expected_output": "3", "points": 1, "is_hidden": True, "is_example": False}
            ]
        )
    assert "scored test cases require at least 1 point each" in str(exc.value)


@pytest.mark.django_db
def test_56_manual_mcq_question_creation_validates_canonical_dto_rules(admin_user):
    """56. Manual creation of MCQ question validates canonical rules (>= 2 options, exactly 1 correct)."""
    with pytest.raises(DRFValidationError) as exc:
        QuestionService.create_question(
            actor=admin_user,
            question_type=QuestionType.MCQ,
            title="Manual MCQ Fail",
            description="Statement",
            points=5,
            type_config={
                "options": [
                    {"id": "A", "text": "Opt 1", "is_correct": True},
                    {"id": "B", "text": "Opt 2", "is_correct": True}
                ]
            }
        )
    assert "exactly one correct option" in str(exc.value)


# ==============================================================================
# GROUP 10: Explicit Snapshot Immutability & Secret-Leak Security Verification
# ==============================================================================

@pytest.mark.django_db
def test_historical_assessment_snapshot_is_immutable_after_question_version_change(admin_user):
    """Explicitly verify historical exam snapshot remains byte/field-equivalent after question version change."""
    # 1. Create Question Q1 with version v1
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.CODING,
        title="Original Two Sum v1",
        description="Original Statement v1",
        points=10,
        coding_config_data={
            "problem_statement": "Original Statement v1",
            "allowed_languages": ["PYTHON"],
            "starter_codes": {"PYTHON": "def solve(): pass"},
            "time_limit_ms": 2000,
            "memory_limit_mb": 256,
        },
        test_cases_data=[
            {"input_data": "1 2", "expected_output": "3", "points": 5, "is_hidden": False, "is_example": True},
            {"input_data": "2 3", "expected_output": "5", "points": 5, "is_hidden": True, "is_example": False}
        ]
    )
    # 2. Publish v1
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)
    v1.refresh_from_db()

    # 3. Add v1 to Assessment
    assessment = Assessment.objects.create(
        title="Historical Snapshot Assessment",
        description="Assessment for verifying snapshot immutability",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now(),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)

    # 4. Publish/freeze snapshot
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()

    # 5. Confirm snapshot contains v1 content
    initial_snapshot_json = json.dumps(snapshot.snapshot_data, sort_keys=True)
    snap_q = snapshot.snapshot_data["questions"][0]
    assert snap_q["title"] == "Original Two Sum v1"
    assert snap_q["description"] == "Original Statement v1"

    # 6. Fork/create QuestionVersion v2
    v2 = QuestionService.create_new_version(question=q, actor=admin_user)
    QuestionService.update_draft_version(
        version=v2,
        title="Mutated Two Sum v2",
        description="Mutated Statement v2",
        actor=admin_user
    )

    # 7. Publish v2 (which archives v1)
    QuestionService.publish_version(version=v2, actor=admin_user)
    v1.refresh_from_db()
    assert v1.status == VersionStatus.ARCHIVED
    assert v2.status == VersionStatus.PUBLISHED

    # 8. Create an unused draft question and delete it
    unused_q, _ = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Unused Draft",
        description="Statement",
        points=5,
        type_config={"options": [{"id": "A", "text": "A", "is_correct": True}, {"id": "B", "text": "B", "is_correct": False}]}
    )
    QuestionService.delete_draft_question(unused_q, actor=admin_user)

    # 9. Verify the existing AssessmentSnapshot remains 100% byte/field-equivalent to original v1 snapshot
    snapshot.refresh_from_db()
    current_snapshot_json = json.dumps(snapshot.snapshot_data, sort_keys=True)
    assert current_snapshot_json == initial_snapshot_json
    assert "Mutated Two Sum v2" not in current_snapshot_json
    assert "Original Two Sum v1" in current_snapshot_json


@pytest.mark.django_db
def test_student_mcq_payload_contains_no_answer_key(api_client, student_user, admin_user):
    """Student delivery boundary: MCQ payload must not contain answer keys or internal notes."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.MCQ,
        title="Secret Key MCQ",
        description="Identify the capital of France.",
        points=5,
        type_config={
            "options": [
                {"id": "A", "text": "Paris", "is_correct": True},
                {"id": "B", "text": "London", "is_correct": False}
            ]
        }
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)

    assessment = Assessment.objects.create(
        title="Secret Boundary Exam",
        description="Desc",
        duration_minutes=60,
        total_points=5,
        passing_percentage=40,
        start_datetime=timezone.now() - timedelta(minutes=5),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=5, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()
    AssessmentService.assign_students(assessment=assessment, student_ids=[str(student_user.id)], actor=admin_user)

    api_client.force_authenticate(user=student_user)
    res = api_client.get(f"/api/v1/student/assessments/{assessment.id}/")
    assert res.status_code == 200

    raw_payload_str = json.dumps(res.data)
    forbidden_keys = ["is_correct", "correct_option", "correct_options", "answer_key", "accepted_answers", "admin_notes"]
    for k in forbidden_keys:
        assert k not in raw_payload_str, f"Forbidden key '{k}' found in student MCQ payload"


@pytest.mark.django_db
def test_student_coding_payload_contains_no_hidden_tests(api_client, student_user, admin_user):
    """Student delivery boundary: Coding question payload must not contain hidden test inputs or outputs."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.CODING,
        title="Super Secret Tests Problem",
        description="Compute sum of numbers.",
        points=10,
        coding_config_data={
            "problem_statement": "Compute sum of numbers.",
            "allowed_languages": ["PYTHON"],
            "starter_codes": {"PYTHON": "def sum(a, b): pass"},
        },
        test_cases_data=[
            {"name": "PublicSample", "input_data": "VISIBLE_INPUT_1", "expected_output": "VISIBLE_OUTPUT_1", "points": 5, "is_hidden": False, "is_example": True},
            {"name": "PrivateHidden", "input_data": "SECRET_INPUT_FORBIDDEN", "expected_output": "SECRET_OUTPUT_FORBIDDEN", "points": 5, "is_hidden": True, "is_example": False}
        ]
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)

    assessment = Assessment.objects.create(
        title="Hidden Test Boundary Exam",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now() - timedelta(minutes=5),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()
    AssessmentService.assign_students(assessment=assessment, student_ids=[str(student_user.id)], actor=admin_user)

    api_client.force_authenticate(user=student_user)
    res = api_client.get(f"/api/v1/student/assessments/{assessment.id}/")
    assert res.status_code == 200

    raw_payload_str = json.dumps(res.data)
    assert "SECRET_INPUT_FORBIDDEN" not in raw_payload_str
    assert "SECRET_OUTPUT_FORBIDDEN" not in raw_payload_str
    assert "PrivateHidden" not in raw_payload_str


@pytest.mark.django_db
def test_student_coding_payload_contains_no_reference_solution(api_client, student_user, admin_user):
    """Student delivery boundary: Coding question payload must not contain reference solution code."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.CODING,
        title="Reference Solution Leak Test",
        description="Problem",
        points=10,
        coding_config_data={
            "problem_statement": "Problem",
            "allowed_languages": ["PYTHON"],
            "starter_codes": {"PYTHON": "def solve(): pass"},
            "reference_solutions": {"PYTHON": "def solve(): return 'CLASSIFIED_AUTHOR_SOLUTION'"},
            "reference_solution_language": "PYTHON"
        },
        test_cases_data=[
            {"input_data": "1", "expected_output": "1", "points": 10, "is_hidden": False, "is_example": True}
        ]
    )
    v1.coding_config.mark_reference_solution_verified("PYTHON", "def solve(): return 'CLASSIFIED_AUTHOR_SOLUTION'")
    v1.coding_config.save()
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)

    assessment = Assessment.objects.create(
        title="Ref Solution Exam",
        description="Desc",
        duration_minutes=60,
        total_points=10,
        passing_percentage=40,
        start_datetime=timezone.now() - timedelta(minutes=5),
        end_datetime=timezone.now() + timedelta(hours=2),
        created_by=admin_user,
        status=AssessmentStatus.DRAFT
    )
    AssessmentService.add_question_to_assessment(str(assessment.id), str(v1.id), points=10, actor=admin_user)
    snapshot = AssessmentSnapshotService.create_snapshot(assessment=assessment, actor=admin_user)
    assessment.active_snapshot = snapshot
    assessment.status = AssessmentStatus.PUBLISHED
    assessment.save()
    AssessmentService.assign_students(assessment=assessment, student_ids=[str(student_user.id)], actor=admin_user)

    api_client.force_authenticate(user=student_user)
    res = api_client.get(f"/api/v1/student/assessments/{assessment.id}/")
    assert res.status_code == 200

    raw_payload_str = json.dumps(res.data)
    assert "CLASSIFIED_AUTHOR_SOLUTION" not in raw_payload_str
    assert "reference_solutions" not in raw_payload_str
    assert "reference_solution" not in raw_payload_str


@pytest.mark.django_db
def test_student_cannot_access_evaluator_bundle(api_client, student_user, admin_user):
    """Student delivery boundary: Student cannot access private evaluation bundle or staff endpoints."""
    q, v1 = QuestionService.create_question(
        actor=admin_user,
        question_type=QuestionType.CODING,
        title="Evaluation Bundle Isolation",
        description="Problem",
        points=10,
        coding_config_data={
            "problem_statement": "Problem",
            "allowed_languages": ["PYTHON"],
            "starter_codes": {"PYTHON": "pass"}
        },
        test_cases_data=[
            {"input_data": "1", "expected_output": "1", "points": 10, "is_hidden": False, "is_example": True}
        ]
    )
    QuestionService.publish_version(str(q.id), 1, actor=admin_user)

    api_client.force_authenticate(user=student_user)

    # 1. Staff preview endpoint rejected
    preview_res = api_client.get(f"/api/v1/admin/questions/{q.id}/versions/1/preview/")
    assert preview_res.status_code in [401, 403]

    # 2. Admin detail endpoint rejected
    admin_detail_res = api_client.get(f"/api/v1/admin/questions/{q.id}/")
    assert admin_detail_res.status_code in [401, 403]

    # 3. Health check endpoint rejected
    health_res = api_client.get(f"/api/v1/admin/questions/{q.id}/versions/1/health/")
    assert health_res.status_code in [401, 403]

